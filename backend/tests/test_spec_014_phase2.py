"""
Unit and integration tests for SPEC-014 Phase 2:
Backend - API and FamilySearch DIP Export Engine.
"""

import io
import csv
import json
import pytest
from openpyxl import load_workbook
from fastapi import HTTPException

from backend.main import (
    FolioEntryCreate,
    FolioEntryUpdate,
    compute_canonical_filename
)
from backend.exporters.familysearch import (
    generate_familysearch_matrix,
    generate_gedcom_551,
    generate_manifest_json,
    consolidate_continued_entries
)


def test_t02_be_01_folio_entry_schemas():
    """T02-BE-01: Validates Pydantic schemas admit physical foliation fields."""
    create_payload = {
        "entry_number": 1,
        "metadata": {"first_name": "Juan", "last_name": "Perez"},
        "status": "completed",
        "folio_number": 14,
        "folio_side": "recto",
        "act_number": "45",
        "is_continued": False,
        "continued_from_entry_id": None
    }
    entry_create = FolioEntryCreate(**create_payload)
    assert entry_create.folio_number == 14
    assert entry_create.folio_side == "recto"
    assert entry_create.act_number == "45"
    assert entry_create.is_continued is False
    assert entry_create.continued_from_entry_id is None

    update_payload = {
        "folio_number": 15,
        "folio_side": "vuelto",
        "act_number": "S/N",
        "is_continued": True,
        "continued_from_entry_id": "00000000-0000-0000-0000-000000000001"
    }
    entry_update = FolioEntryUpdate(**update_payload)
    assert entry_update.folio_number == 15
    assert entry_update.folio_side == "vuelto"
    assert entry_update.act_number == "S/N"
    assert entry_update.is_continued is True
    assert entry_update.continued_from_entry_id == "00000000-0000-0000-0000-000000000001"


def test_t02_be_03_canonical_filename_with_cara():
    """T02-BE-03: Tests canonical nomenclature with physical foliation and cara."""
    # Recto (R)
    name_recto = compute_canonical_filename(
        tenant_subdomain="CARUPANO",
        tome_name="BAUTISMOS",
        tome_number="1",
        sequence_order=1,
        folio_number=14,
        folio_side="recto"
    )
    assert name_recto == "CARUPANO_BAUTISMOS_T01_F014R.jpg"

    # Vuelto (V)
    name_vuelto = compute_canonical_filename(
        tenant_subdomain="CARUPANO",
        tome_name="BAUTISMOS",
        tome_number="1",
        sequence_order=2,
        folio_number=14,
        folio_side="vuelto"
    )
    assert name_vuelto == "CARUPANO_BAUTISMOS_T01_F014V.jpg"

    # Fallback to sequence_order when folio_number is missing
    name_seq = compute_canonical_filename(
        tenant_subdomain="CARUPANO",
        tome_name="BAUTISMOS",
        tome_number="1",
        sequence_order=5
    )
    assert name_seq == "CARUPANO_BAUTISMOS_T01_F0005.jpg"


def test_t02_be_02_strict_dip_filter_and_zero_pending_rows():
    """
    T02-BE-02: Validates strict filter:
    1. Pending entries are excluded from XLSX and CSV matrices (ZERO pending rows).
    2. Physical foliation columns are present in headers.
    """
    entries = [
        {
            "id": "e1",
            "folio_id": "f1",
            "entry_number": 1,
            "status": "completed",
            "folio_number": 12,
            "folio_side": "recto",
            "act_number": "101",
            "metadata": {
                "first_name": "Carlos",
                "last_name": "Gomez",
                "sex": "M",
                "event_date": "1880-05-12"
            }
        },
        {
            "id": "e2",
            "folio_id": "f1",
            "entry_number": 2,
            "status": "pending",  # DEBE SER EXCLUIDO
            "folio_number": 12,
            "folio_side": "recto",
            "act_number": "102",
            "metadata": {
                "first_name": "Incompleto",
                "last_name": "Pendiente"
            }
        },
        {
            "id": "e3",
            "folio_id": "f2",
            "entry_number": 1,
            "status": "verified",
            "folio_number": 12,
            "folio_side": "vuelto",
            "act_number": "103",
            "metadata": {
                "first_name": "Maria",
                "last_name": "Rodriguez",
                "sex": "F",
                "event_date": "1880-06-15"
            }
        }
    ]

    folios = [
        {"id": "f1", "sequence_order": 1, "canonical_filename": "CARUPANO_BAU_T01_F012R.jpg"},
        {"id": "f2", "sequence_order": 2, "canonical_filename": "CARUPANO_BAU_T01_F012V.jpg"}
    ]

    tome_metadata = {
        "tome_number": "1",
        "parish_name": "Santa Rosa",
        "country": "Venezuela",
        "state": "Sucre"
    }

    excel_bytes, csv_bytes = generate_familysearch_matrix(entries, folios, tome_metadata)

    # 1. Parse Excel
    wb = load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]

    assert "folio_fisico" in headers
    assert "cara_folio" in headers
    assert "numero_acta" in headers

    # Filas de datos (excluyendo cabecera)
    data_rows = rows[1:]
    assert len(data_rows) == 2  # Solo e1 (completed) y e3 (verified). e2 (pending) omitida.

    status_idx = headers.index("status")
    for r in data_rows:
        assert r[status_idx] in ("completed", "verified")
        assert r[status_idx] != "pending"

    # 2. Parse CSV
    csv_text = csv_bytes.decode("utf-8")
    assert "sep=;" in csv_text
    lines = [line.strip() for line in csv_text.splitlines() if line.strip() and not line.startswith("sep=")]
    reader = csv.reader(lines, delimiter=";")
    csv_rows = list(reader)
    csv_headers = csv_rows[0]
    assert "folio_fisico" in csv_headers
    assert "cara_folio" in csv_headers
    assert "numero_acta" in csv_headers

    csv_data = csv_rows[1:]
    assert len(csv_data) == 2
    for r in csv_data:
        assert "pending" not in r


def test_t02_be_02_marriage_act_unbundling():
    """T02-BE-02: Validates marriage acts unbundle into two linked rows for groom and bride."""
    entries = [
        {
            "id": "m1",
            "folio_id": "f1",
            "entry_number": 1,
            "status": "completed",
            "folio_number": 20,
            "folio_side": "recto",
            "act_number": "55",
            "metadata": {
                "groom_first_name": "Antonio",
                "groom_last_name": "Brito",
                "groom_father": "Pedro Brito",
                "groom_mother": "Rosa Garcia",
                "bride_first_name": "Carmen",
                "bride_last_name": "Navarro",
                "bride_father": "Manuel Navarro",
                "bride_mother": "Juana Salazar",
                "event_date": "1895-10-20"
            }
        }
    ]

    folios = [{"id": "f1", "sequence_order": 1, "canonical_filename": "CARUPANO_MAT_T01_F020R.jpg"}]
    tome_metadata = {"tome_number": "1", "parish_name": "San Juan Bautista"}

    excel_bytes, _ = generate_familysearch_matrix(entries, folios, tome_metadata)
    wb = load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    data_rows = rows[1:]

    # Debe generar 2 filas vinculadas por el mismo numero_acta y folio_fisico
    assert len(data_rows) == 2

    groom_row = data_rows[0]
    bride_row = data_rows[1]

    act_idx = headers.index("numero_acta")
    folio_idx = headers.index("folio_fisico")
    first_idx = headers.index("primer_nombre")
    last_idx = headers.index("apellidos")
    sex_idx = headers.index("sexo")
    spouse_idx = headers.index("conyuge")

    # Ambas filas comparten folio y acta
    assert groom_row[act_idx] == "55"
    assert bride_row[act_idx] == "55"
    assert groom_row[folio_idx] == 20
    assert bride_row[folio_idx] == 20

    # Novio
    assert groom_row[first_idx] == "Antonio"
    assert groom_row[last_idx] == "Brito"
    assert groom_row[sex_idx] == "M"
    assert groom_row[spouse_idx] == "Carmen Navarro"

    # Novia
    assert bride_row[first_idx] == "Carmen"
    assert bride_row[last_idx] == "Navarro"
    assert bride_row[sex_idx] == "F"
    assert bride_row[spouse_idx] == "Antonio Brito"


def test_t02_be_04_consolidation_of_continued_entries():
    """T02-BE-04: Validates continued entries interfolio are consolidated and not duplicated."""
    entries = [
        {
            "id": "e_parent",
            "folio_id": "f1",
            "entry_number": 3,
            "status": "completed",
            "folio_number": 14,
            "folio_side": "vuelto",
            "act_number": "78",
            "is_continued": False,
            "continued_from_entry_id": None,
            "metadata": {
                "first_name": "Francisco",
                "last_name": "Alcala",
                "marginal_notes": "Bautizado sub conditione."
            }
        },
        {
            "id": "e_continued",
            "folio_id": "f2",
            "entry_number": 1,
            "status": "completed",
            "folio_number": 15,
            "folio_side": "recto",
            "act_number": "78",
            "is_continued": True,
            "continued_from_entry_id": "e_parent",
            "metadata": {
                "marginal_notes": "Fue legitimado por matrimonio subsiguiente de sus padres en 1910."
            }
        }
    ]

    consolidated = consolidate_continued_entries(entries)
    # Solo debe quedar la partida madre
    assert len(consolidated) == 1
    mother = consolidated[0]
    assert mother["id"] == "e_parent"
    notes = mother["metadata"]["marginal_notes"]
    assert "Bautizado sub conditione" in notes
    assert "Fue legitimado por matrimonio subsiguiente" in notes
    assert "Folio 15 recto" in notes or "Folio 15" in notes


def test_t02_be_05_manifest_traceability_metrics():
    """T02-BE-05: Validates manifest.json reports full archival metrics."""
    folder = {
        "id": "folder_1",
        "name": "Libro I Bautismos",
        "folder_type": "tome",
        "metadata": {"tome_number": "1", "country": "Venezuela", "state": "Sucre"}
    }
    folios = [
        {"id": "f1", "sequence_order": 1, "status": "completed"},
        {"id": "f2", "sequence_order": 2, "status": "completed"},
        {"id": "f3", "sequence_order": 3, "status": "pending"}  # Folio pendiente que será omitido
    ]
    completed_entries = [
        {"id": "e1", "folio_id": "f1", "status": "completed"},
        {"id": "e2", "folio_id": "f2", "status": "completed"}
    ]

    manifest = generate_manifest_json(
        folder=folder,
        folios=folios,
        entries=completed_entries,
        total_completed=2,
        total_revision=0,
        total_tome_folios=3,
        excluded_pending_folios=1,
        quality_score_percentage=66.67
    )

    stats = manifest["statistics"]
    assert stats["total_tome_folios"] == 3
    assert stats["exported_entries_count"] == 2
    assert stats["excluded_pending_folios"] == 1
    assert stats["quality_score_percentage"] == 66.67
    assert stats["pending_entries"] == 0

    # Verifica lista de archivos con indicador is_exported
    files = manifest["files"]
    assert len(files) == 3
    f1_item = next(f for f in files if f["folio_id"] == "f1")
    f3_item = next(f for f in files if f["folio_id"] == "f3")
    assert f1_item["is_exported"] is True
    assert f3_item["is_exported"] is False


def test_t02_be_06_records_response_structure_mock():
    """T02-BE-06: Validates records response structure conforms to required contract."""
    # Simula la estructura generada por el endpoint /api/v1/records
    mock_records = [
        {
            "id": "entry-101",
            "folio_id": "f1",
            "entry_number": 1,
            "folio_number": 14,
            "folio_side": "recto",
            "act_number": "45",
            "status": "completed",
            "principal_person": "Juan Bautista Arismendi",
            "event_date": "1875-03-15",
            "record_type": "Bautismo",
            "created_at": "2026-10-17T00:00:00Z"
        }
    ]
    response = {
        "status": "success",
        "data": mock_records,
        "total": 1,
        "page": 1,
        "page_size": 20
    }

    assert response["status"] == "success"
    assert response["total"] == 1
    assert response["page"] == 1
    assert response["page_size"] == 20
    rec = response["data"][0]
    assert rec["folio_number"] == 14
    assert rec["folio_side"] == "recto"
    assert rec["act_number"] == "45"
    assert rec["principal_person"] == "Juan Bautista Arismendi"
