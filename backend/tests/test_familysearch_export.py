import pytest
import io
import json
import zipfile
import openpyxl
from backend.exporters.familysearch import (
    generate_familysearch_matrix,
    generate_gedcom_551,
    generate_manifest_json,
    parse_gedcom_date
)

def test_parse_gedcom_date():
    assert parse_gedcom_date("1855-04-12") == "12 APR 1855"
    assert parse_gedcom_date("1855-04") == "APR 1855"
    assert parse_gedcom_date("1855") == "1855"
    assert parse_gedcom_date("15/04/1855") == "15 APR 1855"
    assert parse_gedcom_date(None) is None
    assert parse_gedcom_date("") is None


def test_familysearch_matrix_generation():
    folios = [
        {
            "id": "folio-1",
            "sequence_order": 1,
            "canonical_filename": "DIOC_CAR_BAU_T01_F0001.jpg",
            "original_filename": "IMG_001.jpg"
        }
    ]
    entries = [
        {
            "folio_id": "folio-1",
            "entry_number": 1,
            "status": "completed",
            "metadata": {
                "primer_nombre": "Juan",
                "apellidos": "Pérez Gómez",
                "sexo": "M",
                "fecha_acta": "1855-04-12",
                "fecha_nacimiento": "1855-04-10",
                "padre": "Pedro Pérez",
                "madre": "María Gómez",
                "padrinos": "José García y Ana Romero",
                "observaciones": "Acta legible"
            }
        }
    ]
    tome_metadata = {
        "parish_name": "Parroquia Santa Rosa",
        "tome_number": "1",
        "diocese_name": "Diócesis de Carúpano"
    }

    excel_bytes, csv_bytes = generate_familysearch_matrix(entries, folios, tome_metadata)

    # 1. Validar Excel
    assert len(excel_bytes) > 0
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.title == "Matriz FamilySearch"
    headers = [cell.value for cell in ws[1]]
    assert "image_file_name" in headers
    assert "entry_number" in headers
    assert "primer_nombre" in headers

    row_2 = [cell.value for cell in ws[2]]
    assert row_2[0] == "DIOC_CAR_BAU_T01_F0001.jpg"
    assert row_2[1] == 1  # folio_number / sequence_order
    assert row_2[2] == 1  # entry_number
    assert row_2[4] == "Juan"
    assert row_2[5] == "Pérez Gómez"

    # 2. Validar CSV
    assert csv_bytes.startswith(b"\xef\xbb\xbf")  # BOM UTF-8
    csv_str = csv_bytes.decode("utf-8")
    assert csv_str.startswith("\ufeffsep=;\r\n")
    lines = csv_str.splitlines()
    assert "image_file_name" in lines[1]
    assert "DIOC_CAR_BAU_T01_F0001.jpg;1;1;completed;Juan;Pérez Gómez" in lines[2]


def test_gedcom_551_generation():
    folios = [
        {
            "id": "folio-1",
            "sequence_order": 42,
            "canonical_filename": "DIOC_CAR_BAU_T01_F0042.jpg"
        }
    ]
    entries = [
        {
            "folio_id": "folio-1",
            "entry_number": 1,
            "status": "completed",
            "metadata": {
                "primer_nombre": "Juan",
                "apellidos": "Pérez",
                "sexo": "M",
                "fecha_acta": "1855-04-12",
                "padre": "Pedro Pérez",
                "madre": "María González",
                "padrinos": "José Martínez"
            }
        }
    ]
    tome_metadata = {
        "parish_name": "Parroquia San José",
        "diocese_name": "Diócesis de Carúpano",
        "series_code": "BAU"
    }

    gedcom_str = generate_gedcom_551(entries, folios, tome_metadata)

    # Validaciones obligatorias de la especificación GEDCOM 5.5.1
    assert "0 HEAD" in gedcom_str
    assert "1 SOUR CODEXFLOW" in gedcom_str
    assert "1 GEDC" in gedcom_str
    assert "2 VERS 5.5.1" in gedcom_str
    assert "2 FORM LINEAGE-LINKED" in gedcom_str
    assert "1 CHAR UTF-8" in gedcom_str

    # Individuo principal
    assert "0 @I1@ INDI" in gedcom_str
    assert "1 NAME Juan /Pérez/" in gedcom_str
    assert "1 SEX M" in gedcom_str
    assert "1 CHR" in gedcom_str
    assert "2 DATE 12 APR 1855" in gedcom_str
    assert "Parroquia San José, Diócesis de Carúpano" in gedcom_str

    # Objeto digital asociado
    assert "1 OBJE" in gedcom_str
    assert "2 FILE images/DIOC_CAR_BAU_T01_F0042.jpg" in gedcom_str

    # Padres y familia
    assert "0 @I2@ INDI" in gedcom_str
    assert "1 NAME Pedro Pérez //" in gedcom_str
    assert "0 @I3@ INDI" in gedcom_str
    assert "1 NAME María González //" in gedcom_str
    assert "0 @F1@ FAM" in gedcom_str
    assert "1 HUSB @I2@" in gedcom_str
    assert "1 WIFE @I3@" in gedcom_str
    assert "1 CHIL @I1@" in gedcom_str

    # Cierre
    assert "0 TRLR" in gedcom_str


def test_manifest_json_generation():
    folder = {
        "id": "tome-001",
        "name": "Libro 01 (1850-1875)",
        "folder_type": "tome",
        "metadata": {
            "tome_number": "1",
            "year_start": 1850,
            "year_end": 1875,
            "parish_name": "San José",
            "diocese_name": "Carúpano"
        }
    }
    folios = [
        {
            "id": "f-1",
            "sequence_order": 1,
            "original_filename": "IMG_001.jpg",
            "canonical_filename": "DIOC_CAR_BAU_T01_F0001.jpg",
            "status": "completed"
        }
    ]
    entries = [
        {
            "folio_id": "f-1",
            "entry_number": 1,
            "status": "completed",
            "metadata": {"primer_nombre": "Juan"}
        },
        {
            "folio_id": "f-1",
            "entry_number": 2,
            "status": "revision",
            "metadata": {"primer_nombre": "Pedro"}
        }
    ]

    manifest = generate_manifest_json(folder, folios, entries, total_completed=1, total_revision=1)

    assert manifest["project_name"] == "CodexFlow - FamilySearch DIP Export"
    assert manifest["dip_version"] == "1.0"
    assert manifest["standard"] == "FamilySearch DIP Homologation"
    assert "generated_at" in manifest
    assert manifest["tome_info"]["name"] == "Libro 01 (1850-1875)"
    assert manifest["tome_info"]["tome_number"] == "1"
    assert "geographic_coverage" in manifest["tome_info"]
    assert manifest["tome_info"]["geographic_coverage"]["is_partially_inherited"] is True
    assert manifest["statistics"]["total_folios"] == 1
    assert manifest["statistics"]["total_entries"] == 2
    assert manifest["statistics"]["completed_entries"] == 1
    assert manifest["statistics"]["revision_entries"] == 1
    assert manifest["statistics"]["quality_score_percentage"] == 50.0
    assert len(manifest["files"]) == 1
    assert manifest["files"][0]["canonical_filename"] == "DIOC_CAR_BAU_T01_F0001.jpg"
    assert manifest["files"][0]["entries_count"] == 2


def test_familysearch_matrix_with_geographic_hierarchy():
    folios = [
        {
            "id": "folio-100",
            "sequence_order": 5,
            "canonical_filename": "VEN_SUC_BER_SRO_BAU_T01_F0005.jpg",
            "original_filename": "folio_5.jpg"
        }
    ]
    entries = [
        {
            "folio_id": "folio-100",
            "entry_number": 1,
            "status": "completed",
            "metadata": {
                "primer_nombre": "Carlos",
                "apellidos": "Blanco",
                "sexo": "M",
                "fecha_acta": "1880-05-20"
            }
        },
        {
            "folio_id": "folio-100",
            "entry_number": 2,
            "status": "completed",
            "metadata": {
                "primer_nombre": "Elena",
                "apellidos": "Rojas",
                "sexo": "F",
                "fecha_acta": "1880-05-22",
                "ciudad": "Playa Grande",
                "canonical_place": "Playa Grande, Bermúdez, Sucre, Venezuela"
            }
        }
    ]
    tome_metadata = {
        "parish_name": "Santa Rosa",
        "tome_number": "1",
        "diocese_name": "Carúpano",
        "country": "Venezuela",
        "state": "Sucre",
        "municipality": "Bermúdez",
        "city_or_locality": "Carúpano",
        "canonical_place": "Carúpano, Bermúdez, Sucre, Venezuela"
    }

    excel_bytes, csv_bytes = generate_familysearch_matrix(entries, folios, tome_metadata)

    # 1. Validar Excel
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    for expected_col in ["pais", "estado", "municipio", "ciudad_localidad", "lugar_evento_estandarizado"]:
        assert expected_col in headers

    row_carlos = [cell.value for cell in ws[2]]
    # Columnas 16 a 20 (índices 15 a 19): pais, estado, municipio, ciudad_localidad, lugar_evento_estandarizado
    assert row_carlos[headers.index("pais")] == "Venezuela"
    assert row_carlos[headers.index("estado")] == "Sucre"
    assert row_carlos[headers.index("municipio")] == "Bermúdez"
    assert row_carlos[headers.index("ciudad_localidad")] == "Carúpano"
    assert row_carlos[headers.index("lugar_evento_estandarizado")] == "Carúpano, Bermúdez, Sucre, Venezuela"

    # Fila con sobreescritura individual en Elena
    row_elena = [cell.value for cell in ws[3]]
    assert row_elena[headers.index("ciudad_localidad")] == "Playa Grande"
    assert row_elena[headers.index("lugar_evento_estandarizado")] == "Playa Grande, Bermúdez, Sucre, Venezuela"

    # 2. Validar CSV
    csv_str = csv_bytes.decode("utf-8")
    assert "pais;estado;municipio;ciudad_localidad;lugar_evento_estandarizado" in csv_str
    assert "Venezuela;Sucre;Bermúdez;Carúpano;Carúpano, Bermúdez, Sucre, Venezuela" in csv_str
    assert "Venezuela;Sucre;Bermúdez;Playa Grande;Playa Grande, Bermúdez, Sucre, Venezuela" in csv_str


def test_gedcom_551_plac_inclusion_and_omission():
    folios = [{"id": "f1", "sequence_order": 1, "canonical_filename": "folio_1.jpg"}]
    entries = [
        {
            "folio_id": "f1",
            "entry_number": 1,
            "status": "completed",
            "metadata": {
                "primer_nombre": "Luisa",
                "apellidos": "Mendoza",
                "fecha_acta": "1890-01-15",
                "fecha_nacimiento": "1890-01-10"
            }
        }
    ]

    # Caso A: Con canonical_place presente
    tome_meta_with_geo = {
        "parish_name": "San José",
        "canonical_place": "Carúpano, Bermúdez, Sucre, Venezuela"
    }
    gedcom_with_geo = generate_gedcom_551(entries, folios, tome_meta_with_geo)
    assert "2 PLAC Carúpano, Bermúdez, Sucre, Venezuela" in gedcom_with_geo

    # Caso B: Sin datos geográficos en ningún nivel
    tome_meta_empty = {}
    gedcom_empty = generate_gedcom_551(entries, folios, tome_meta_empty)
    assert "2 PLAC" not in gedcom_empty
    # Verificar que el bloque INDI sigue intacto y válido
    assert "0 @I1@ INDI" in gedcom_empty
    assert "1 NAME Luisa /Mendoza/" in gedcom_empty
    assert "1 CHR" in gedcom_empty
    assert "2 DATE 15 JAN 1890" in gedcom_empty
    assert "0 TRLR" in gedcom_empty


def test_manifest_json_geographic_coverage():
    folder = {
        "id": "tome-xyz",
        "name": "Libro de Bautismos 1890",
        "folder_type": "tome",
        "metadata": {
            "tome_number": "3"
        }
    }
    tome_metadata = {
        "country": "Venezuela",
        "state": "Sucre",
        "municipality": "Bermúdez",
        "city_or_locality": "Carúpano",
        "canonical_place": "Carúpano, Bermúdez, Sucre, Venezuela",
        "parish_name": "Santa Rosa",
        "diocese_name": "Diócesis de Carúpano"
    }
    folios = [{"id": "f1", "sequence_order": 1, "canonical_filename": "f1.jpg", "status": "completed"}]
    entries = [{"folio_id": "f1", "entry_number": 1, "status": "completed", "metadata": {}}]

    manifest = generate_manifest_json(folder, folios, entries, 1, 0, tome_metadata=tome_metadata)

    assert "geographic_coverage" in manifest["tome_info"]
    geo = manifest["tome_info"]["geographic_coverage"]
    assert geo["country"] == "Venezuela"
    assert geo["state"] == "Sucre"
    assert geo["municipality"] == "Bermúdez"
    assert geo["city_or_locality"] == "Carúpano"
    assert geo["standardized_place"] == "Carúpano, Bermúdez, Sucre, Venezuela"
    assert geo["is_partially_inherited"] is True

